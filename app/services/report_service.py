"""Mesačný prehľad hodín (podklad pre mzdy).

Pre každého zamestnanca spočíta za mesiac: počet smien, odpracované
hodiny, z toho nočné (22:00-06:00), sobotné, nedeľné a hodiny počas
štátneho sviatku, a pracovné dni absencií (dovolenka, PN, OČR, náhradné
voľno).

Pravidlá výpočtu (uvádzajú sa aj v exporte):
  - smena patrí do mesiaca podľa DÁTUMU ZAČIATKU (nočná smena z 30. na
    1. sa počíta celá do mesiaca, v ktorom začala),
  - hodiny sú dĺžka smeny bez odpočítania prestávok,
  - nočné/sobotné/nedeľné/sviatočné hodiny sa určujú podľa skutočných
    kalendárnych dní, cez ktoré smena prechádza (nočná z piatka na
    sobotu má sobotné hodiny po polnoci); sobota/nedeľa a sviatok sa
    môžu prekrývať,
  - pracovný deň = pondelok až piatok mimo štátnych sviatkov (SR).
"""

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from app.services.absence_service import get_absences_in_range
from app.services.calendar_service import month_bounds
from app.services.employee_service import get_employees
from app.services.holiday_service import (
    is_public_holiday,
    working_days_between,
)
from app.services.publish_service import get_draft_shift_ids
from app.services.shift_service import get_shifts_in_range

ABSENCE_COLUMNS = ["Dovolenka", "PN", "OČR", "Náhradné voľno"]

WEEKDAY_SHORT = ["Po", "Ut", "St", "Št", "Pi", "So", "Ne"]


def _overlap_hours(start, end, window_start, window_end):
    latest_start = max(start, window_start)
    earliest_end = min(end, window_end)

    if earliest_end <= latest_start:
        return 0.0

    return (earliest_end - latest_start).total_seconds() / 3600


def split_shift_hours(shift_date, start_time, end_time):
    """Rozdelí smenu na hodiny: ``total``, ``night`` (22-06),
    ``saturday``, ``sunday``, ``holiday``.

    Ak je koniec skorší/rovnaký ako začiatok, smena prechádza cez polnoc.
    """

    day = datetime.strptime(shift_date, "%Y-%m-%d")

    start_hour, start_minute = (int(x) for x in start_time.split(":")[:2])
    end_hour, end_minute = (int(x) for x in end_time.split(":")[:2])

    start = day.replace(hour=start_hour, minute=start_minute)
    end = day.replace(hour=end_hour, minute=end_minute)

    if end <= start:
        end += timedelta(days=1)

    result = {
        "total": (end - start).total_seconds() / 3600,
        "night": 0.0,
        "saturday": 0.0,
        "sunday": 0.0,
        "holiday": 0.0,
    }

    current_day = start.replace(hour=0, minute=0)

    while current_day < end:
        next_day = current_day + timedelta(days=1)
        segment_start = max(start, current_day)
        segment_end = min(end, next_day)
        hours = (segment_end - segment_start).total_seconds() / 3600

        if hours > 0:
            result["night"] += _overlap_hours(
                segment_start, segment_end,
                current_day, current_day.replace(hour=6),
            ) + _overlap_hours(
                segment_start, segment_end,
                current_day.replace(hour=22), next_day,
            )

            weekday = current_day.weekday()

            if weekday == 5:
                result["saturday"] += hours
            elif weekday == 6:
                result["sunday"] += hours

            if is_public_holiday(current_day.date()):
                result["holiday"] += hours

        current_day = next_day

    return result


@dataclass
class EmployeeHours:
    employee_id: int
    name: str
    branch_name: str
    shifts: int = 0
    hours: float = 0.0
    night: float = 0.0
    saturday: float = 0.0
    sunday: float = 0.0
    holiday: float = 0.0
    absence_days: dict = field(
        default_factory=lambda: {name: 0 for name in ABSENCE_COLUMNS}
    )


@dataclass
class HoursReport:
    year: int
    month: int
    working_days: int
    rows: list
    totals: EmployeeHours
    shift_details: list
    only_published: bool


def build_hours_report(year, month, branch_ids=None, only_published=False):
    """Zostaví mesačný prehľad hodín pre zamestnancov daných pobočiek
    (``None`` = všetky). Zamestnanci bez smeny a absencie v mesiaci sa
    neuvádzajú."""

    first_iso, last_iso = month_bounds(year, month)
    first_day = date.fromisoformat(first_iso)
    last_day = date.fromisoformat(last_iso)

    employees = {
        row[0]: row for row in get_employees(branch_ids)
    }

    rows = {}

    def row_for(employee_id):
        if employee_id not in rows:
            employee = employees[employee_id]
            rows[employee_id] = EmployeeHours(
                employee_id=employee_id,
                name=f"{employee[1]} {employee[2]}",
                branch_name=employee[7],
            )

        return rows[employee_id]

    shifts = get_shifts_in_range(
        first_iso, last_iso, branch_ids, published_only=only_published
    )
    draft_ids = get_draft_shift_ids(first_iso, last_iso, branch_ids)

    details = []

    for shift in shifts:
        employee_id = shift[7]

        if employee_id not in employees:
            continue

        split = split_shift_hours(shift[3], shift[4], shift[5])
        row = row_for(employee_id)

        row.shifts += 1
        row.hours += split["total"]
        row.night += split["night"]
        row.saturday += split["saturday"]
        row.sunday += split["sunday"]
        row.holiday += split["holiday"]

        shift_day = date.fromisoformat(shift[3])

        details.append(
            {
                "name": row.name,
                "branch_name": row.branch_name,
                "date": shift[3],
                "weekday": WEEKDAY_SHORT[shift_day.weekday()],
                "start": shift[4],
                "end": shift[5],
                "hours": split["total"],
                "night": split["night"],
                "saturday": split["saturday"],
                "sunday": split["sunday"],
                "holiday": split["holiday"],
                "shift_type": shift[6],
                "department": shift[9] or "",
                "status": "koncept" if shift[0] in draft_ids else "zverejnená",
            }
        )

    for absence in get_absences_in_range(first_iso, last_iso, branch_ids):
        employee_id = absence[3]

        if employee_id not in employees:
            continue

        start = max(date.fromisoformat(absence[5]), first_day)
        end = min(date.fromisoformat(absence[6]), last_day)
        days = working_days_between(start, end)

        row = row_for(employee_id)
        key = absence[4] if absence[4] in ABSENCE_COLUMNS else None

        if key:
            row.absence_days[key] += days

    ordered = sorted(
        rows.values(), key=lambda r: (r.name.split(" ", 1)[-1].casefold(), r.name.casefold())
    )

    totals = EmployeeHours(employee_id=0, name="Spolu", branch_name="")

    for row in ordered:
        totals.shifts += row.shifts
        totals.hours += row.hours
        totals.night += row.night
        totals.saturday += row.saturday
        totals.sunday += row.sunday
        totals.holiday += row.holiday

        for name in ABSENCE_COLUMNS:
            totals.absence_days[name] += row.absence_days[name]

    details.sort(key=lambda d: (d["name"].split(" ", 1)[-1].casefold(), d["date"], d["start"]))

    return HoursReport(
        year=year,
        month=month,
        working_days=working_days_between(first_day, last_day),
        rows=ordered,
        totals=totals,
        shift_details=details,
        only_published=only_published,
    )


# ---------------------------------------------------------------------
# Exporty
# ---------------------------------------------------------------------

OVERVIEW_HEADERS = [
    "Zamestnanec",
    "Pobočka",
    "Smien",
    "Hodiny spolu",
    "Nočné (22-06)",
    "Sobota",
    "Nedeľa",
    "Sviatok",
    "Dovolenka (prac. dní)",
    "PN (prac. dní)",
    "OČR (prac. dní)",
    "Náhradné voľno (prac. dní)",
]

DETAIL_HEADERS = [
    "Zamestnanec",
    "Pobočka",
    "Dátum",
    "Deň",
    "Začiatok",
    "Koniec",
    "Hodiny",
    "Nočné (22-06)",
    "Sobota",
    "Nedeľa",
    "Sviatok",
    "Typ smeny",
    "Oddelenie",
    "Stav",
]

NOTES = [
    "Ako sa hodiny počítajú:",
    "Smena patrí do mesiaca podľa dátumu jej začiatku (nočná smena z "
    "posledného dňa mesiaca sa počíta celá do tohto mesiaca).",
    "Hodiny sú dĺžka smeny bez odpočítania prestávok.",
    "Nočné hodiny: 22:00 - 06:00.",
    "Sobotné, nedeľné a sviatočné hodiny sa určujú podľa skutočných "
    "kalendárnych dní, cez ktoré smena prechádza (nočná z piatka na "
    "sobotu má sobotné hodiny po polnoci). Sobota/nedeľa a sviatok sa "
    "môžu prekrývať.",
    "Pracovný deň = pondelok až piatok mimo štátnych sviatkov SR; v "
    "absenciách sa počítajú pracovné dni, ktoré spadajú do mesiaca.",
    "Prehľad je podklad - výsledné mzdové nároky určuje mzdová účtovníčka.",
]


def _overview_row(row):
    return [
        row.name,
        row.branch_name,
        row.shifts,
        round(row.hours, 2),
        round(row.night, 2),
        round(row.saturday, 2),
        round(row.sunday, 2),
        round(row.holiday, 2),
        *(row.absence_days[name] for name in ABSENCE_COLUMNS),
    ]


def build_hours_workbook(report):
    """XLSX: listy Prehľad, Smeny (podklad k hodinám) a Poznámky."""

    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    workbook = Workbook()

    def fill_sheet(sheet, headers, data_rows, bold_last=False):
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.font = header_font
            cell.fill = header_fill

        for row_number, values in enumerate(data_rows, start=2):
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row_number, column=column, value=value)

        if bold_last and data_rows:
            for cell in sheet[len(data_rows) + 1]:
                cell.font = Font(bold=True)

        for index, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(
                len(header) + 2, 14
            )

        sheet.freeze_panes = "A2"

    overview = workbook.active
    overview.title = "Prehľad"
    fill_sheet(
        overview,
        OVERVIEW_HEADERS,
        [_overview_row(r) for r in report.rows] + [_overview_row(report.totals)],
        bold_last=True,
    )

    details = workbook.create_sheet("Smeny")
    fill_sheet(
        details,
        DETAIL_HEADERS,
        [
            [
                d["name"], d["branch_name"], d["date"], d["weekday"],
                d["start"], d["end"], round(d["hours"], 2),
                round(d["night"], 2), round(d["saturday"], 2),
                round(d["sunday"], 2), round(d["holiday"], 2),
                d["shift_type"], d["department"], d["status"],
            ]
            for d in report.shift_details
        ],
    )

    notes = workbook.create_sheet("Poznámky")
    notes.column_dimensions["A"].width = 120
    notes.cell(row=1, column=1, value=f"Prehľad hodín {report.month:02d}/{report.year}").font = Font(bold=True)

    scope_text = (
        "Zahrnuté sú len zverejnené smeny."
        if report.only_published
        else "Zahrnuté sú všetky smeny (aj koncepty)."
    )
    notes.cell(row=2, column=1, value=scope_text)
    notes.cell(row=3, column=1, value=f"Pracovných dní v mesiaci: {report.working_days}")

    for index, line in enumerate(NOTES, start=5):
        notes.cell(row=index, column=1, value=line)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return buffer


def build_hours_csv(report):
    """CSV prehľadu (UTF-8 s BOM, oddeľovač ``;`` - Excel ho otvorí
    správne)."""

    import csv
    import io

    buffer = io.StringIO()
    buffer.write("\ufeff")

    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(OVERVIEW_HEADERS)

    for row in report.rows + [report.totals]:
        writer.writerow(
            [
                f"{value:.2f}".replace(".", ",") if isinstance(value, float) else value
                for value in _overview_row(row)
            ]
        )

    return buffer.getvalue()
