"""Testy mesačného prehľadu hodín, sviatkov a rozdelenia hodín smeny."""

import csv
import io
import re
from datetime import date
from io import BytesIO

import pytest

from app.services.holiday_service import (
    easter_sunday,
    is_public_holiday,
    is_working_day,
    slovak_public_holidays,
    working_days_between,
)
from app.services.report_service import split_shift_hours


# ---------------------------------------------------------------------
# Sviatky a pracovné dni
# ---------------------------------------------------------------------

@pytest.mark.parametrize(
    "year, expected",
    [
        (2024, date(2024, 3, 31)),
        (2025, date(2025, 4, 20)),
        (2026, date(2026, 4, 5)),
        (2027, date(2027, 3, 28)),
        (2030, date(2030, 4, 21)),
    ],
)
def test_easter_sunday(year, expected):
    assert easter_sunday(year) == expected


def test_slovak_public_holidays_2026():
    holidays = slovak_public_holidays(2026)

    assert len(holidays) == 15
    assert holidays[date(2026, 4, 3)] == "Veľký piatok"
    assert holidays[date(2026, 4, 6)] == "Veľkonočný pondelok"

    for month, day in [
        (1, 1), (1, 6), (5, 1), (5, 8), (7, 5), (8, 29), (9, 1), (9, 15),
        (11, 1), (11, 17), (12, 24), (12, 25), (12, 26),
    ]:
        assert date(2026, month, day) in holidays

    assert not is_public_holiday(date(2026, 4, 5))   # Veľkonočná nedeľa nie je v zozname


def test_working_days():
    assert is_working_day(date(2026, 9, 2))          # streda
    assert not is_working_day(date(2026, 9, 1))      # sviatok (utorok)
    assert not is_working_day(date(2026, 9, 5))      # sobota

    assert working_days_between(date(2026, 9, 1), date(2026, 9, 4)) == 3
    assert working_days_between(date(2026, 9, 5), date(2026, 9, 6)) == 0
    assert working_days_between(date(2026, 9, 7), date(2026, 9, 7)) == 1
    assert working_days_between(date(2026, 9, 10), date(2026, 9, 9)) == 0
    assert working_days_between(date(2026, 9, 1), date(2026, 9, 30)) == 20


# ---------------------------------------------------------------------
# Rozdelenie hodín smeny
# ---------------------------------------------------------------------

@pytest.mark.parametrize(
    "shift_date, start, end, total, night, saturday, sunday, holiday",
    [
        ("2026-09-07", "06:00", "14:00", 8, 0, 0, 0, 0),      # pondelok cez deň
        ("2026-09-04", "22:00", "06:00", 8, 8, 6, 0, 0),      # piatok -> sobota
        ("2026-09-05", "22:00", "06:00", 8, 8, 2, 6, 0),      # sobota -> nedeľa
        ("2026-09-06", "22:00", "06:00", 8, 8, 0, 2, 0),      # nedeľa -> pondelok
        ("2026-09-01", "06:00", "14:00", 8, 0, 0, 0, 8),      # sviatok
        ("2026-08-31", "22:00", "06:00", 8, 8, 0, 0, 6),      # noc pred sviatkom
        ("2026-09-01", "22:00", "06:00", 8, 8, 0, 0, 2),      # noc zo sviatku
        ("2026-09-07", "18:00", "23:00", 5, 1, 0, 0, 0),      # večer, 1 h nočná
        ("2026-09-07", "04:00", "12:00", 8, 2, 0, 0, 0),      # skoro ráno, 2 h nočné
        ("2026-09-07", "00:00", "08:00", 8, 6, 0, 0, 0),      # od polnoci
        ("2026-09-07", "06:00", "06:00", 24, 8, 0, 0, 0),     # rovnaký čas = 24 h
        ("2026-09-05", "10:00", "10:30", 0.5, 0, 0.5, 0, 0),  # krátka smena v sobotu
    ],
)
def test_split_shift_hours(shift_date, start, end, total, night, saturday, sunday, holiday):
    result = split_shift_hours(shift_date, start, end)

    assert result == pytest.approx(
        {"total": total, "night": night, "saturday": saturday,
         "sunday": sunday, "holiday": holiday}
    )


# ---------------------------------------------------------------------
# Prehľad
# ---------------------------------------------------------------------

def _csrf(html):
    for tag in re.findall(r"<input[^>]*>", html):
        if 'name="csrf_token"' in tag:
            return re.search(r'value="([^"]+)"', tag).group(1)

    raise AssertionError("CSRF token sa nenašiel")


@pytest.fixture
def hw(tmp_path):
    from app.extensions import db
    from app.web.app import create_app

    app = create_app(database_path=tmp_path / "hours.db")
    app.config.update(TESTING=True, WTF_CSRF_ENABLED=True)

    ids = {}

    with app.app_context():
        db.create_all()

        from app.services.absence_service import add_absence
        from app.services.auth_service import (
            create_admin,
            create_employee_user,
            create_manager,
        )
        from app.services.branch_service import add_branch
        from app.services.employee_service import add_employee
        from app.services.shift_service import add_shift

        ids["a"] = add_branch("Košice")
        ids["b"] = add_branch("Prešov")
        ids["anna"] = add_employee("Anna", "Prvá", "Op", 40, ids["a"])
        ids["cyril"] = add_employee("Cyril", "Tretí", "Op", 40, ids["a"])
        ids["boris"] = add_employee("Boris", "Druhý", "Op", 40, ids["b"])
        ids["dana"] = add_employee("Dana", "Štvrtá", "Op", 40, ids["a"])   # bez aktivity

        def shift(emp, day, start, end, published=True, month="09"):
            add_shift(emp, f"2026-{month}-{day}", start, end, "Ranná", published=published)

        shift(ids["anna"], "01", "06:00", "14:00")      # sviatok
        shift(ids["anna"], "04", "22:00", "06:00")      # piatok -> sobota
        shift(ids["anna"], "05", "22:00", "06:00")      # sobota -> nedeľa
        shift(ids["anna"], "07", "06:00", "14:00")
        shift(ids["anna"], "08", "08:00", "12:00", published=False)   # koncept
        shift(ids["anna"], "30", "22:00", "06:00")      # cez koniec mesiaca
        shift(ids["cyril"], "10", "14:00", "22:00")
        shift(ids["cyril"], "11", "18:00", "23:00")
        shift(ids["boris"], "07", "06:00", "14:00")

        add_absence(ids["anna"], "Dovolenka", "2026-08-28", "2026-09-03")
        add_absence(ids["cyril"], "PN", "2026-09-14", "2026-09-16")
        add_absence(ids["boris"], "Dovolenka", "2026-09-21", "2026-09-25")

        create_admin("admin", "adminpass1")
        create_manager("mgr_a", "managerpass1", [ids["a"]])
        create_employee_user(ids["anna"], "emp_a", "emppass12")

    app.test_ids = ids

    return app


def login(app, username, password):
    client = app.test_client()
    token = _csrf(client.get("/login").get_data(as_text=True))
    client.post(
        "/login",
        data={"csrf_token": token, "username": username, "password": password},
    )
    return client


def report(app, year=2026, month=9, branch_ids=None, only_published=False):
    from app.services.report_service import build_hours_report

    with app.app_context():
        return build_hours_report(year, month, branch_ids, only_published)


def by_name(rep):
    return {row.name: row for row in rep.rows}


def test_report_totals_per_employee(hw):
    rep = report(hw)
    rows = by_name(rep)

    anna = rows["Anna Prvá"]
    assert anna.shifts == 6
    assert anna.hours == pytest.approx(44)
    assert anna.night == pytest.approx(24)
    assert anna.saturday == pytest.approx(8)     # 6 h (pi->so) + 2 h (so)
    assert anna.sunday == pytest.approx(6)
    assert anna.holiday == pytest.approx(8)
    assert anna.absence_days == {"Dovolenka": 2, "PN": 0, "OČR": 0, "Náhradné voľno": 0}

    cyril = rows["Cyril Tretí"]
    assert (cyril.shifts, cyril.hours, cyril.night) == (2, pytest.approx(13), pytest.approx(1))
    assert cyril.absence_days["PN"] == 2        # 14., (15. sviatok), 16.

    boris = rows["Boris Druhý"]
    assert boris.hours == pytest.approx(8) and boris.absence_days["Dovolenka"] == 5

    assert rep.working_days == 20
    assert rep.totals.shifts == 9
    assert rep.totals.hours == pytest.approx(65)
    assert rep.totals.absence_days["Dovolenka"] == 7


def test_employees_without_activity_are_left_out_and_sorted_by_surname(hw):
    names = [row.name for row in report(hw).rows]

    assert "Dana Štvrtá" not in names
    assert names == ["Boris Druhý", "Anna Prvá", "Cyril Tretí"]


def test_only_published_excludes_drafts(hw):
    rows = by_name(report(hw, only_published=True))

    assert rows["Anna Prvá"].shifts == 5
    assert rows["Anna Prvá"].hours == pytest.approx(40)


def test_shift_across_month_end_belongs_to_the_month_it_starts_in(hw):
    october = report(hw, month=10)

    assert "Anna Prvá" not in by_name(october)
    assert october.rows == []


def test_absence_is_clipped_to_the_month(hw):
    # Dovolenka 28.8.-3.9.: v auguste 28. (pi), 31. (po) = 2 prac. dni.
    august = by_name(report(hw, month=8))

    assert august["Anna Prvá"].absence_days["Dovolenka"] == 2
    assert august["Anna Prvá"].shifts == 0


def test_branch_scope(hw):
    ids = hw.test_ids

    only_a = report(hw, branch_ids=[ids["a"]])

    assert [r.name for r in only_a.rows] == ["Anna Prvá", "Cyril Tretí"]
    assert only_a.totals.hours == pytest.approx(57)

    assert [r.name for r in report(hw, branch_ids=[ids["b"]]).rows] == ["Boris Druhý"]
    assert report(hw, branch_ids=[]).rows == []


# ---------------------------------------------------------------------
# Stránka a exporty
# ---------------------------------------------------------------------

def test_page_for_admin_and_manager(hw):
    admin_html = login(hw, "admin", "adminpass1").get("/reports/hours?year=2026&month=9").get_data(as_text=True)

    assert "Prehľad hodín" in admin_html
    assert "Anna Prvá" in admin_html and "Boris Druhý" in admin_html
    assert "65,00" in admin_html                     # súčet za všetky pobočky
    assert "pracovných dní: <strong>20</strong>" in admin_html

    manager_html = login(hw, "mgr_a", "managerpass1").get("/reports/hours?year=2026&month=9").get_data(as_text=True)

    assert "Anna Prvá" in manager_html and "Cyril Tretí" in manager_html
    assert "Boris Druhý" not in manager_html
    assert "57,00" in manager_html


def test_branch_filter_and_foreign_branch_is_ignored(hw):
    ids = hw.test_ids

    admin = login(hw, "admin", "adminpass1")
    html = admin.get(f"/reports/hours?year=2026&month=9&branch_id={ids['b']}").get_data(as_text=True)
    assert "Boris Druhý" in html and "Anna Prvá" not in html

    manager = login(hw, "mgr_a", "managerpass1")
    html = manager.get(f"/reports/hours?year=2026&month=9&branch_id={ids['b']}").get_data(as_text=True)
    assert "Boris Druhý" not in html and "Anna Prvá" in html


def test_only_published_option_and_invalid_parameters(hw):
    admin = login(hw, "admin", "adminpass1")

    html = admin.get("/reports/hours?year=2026&month=9&only_published=1").get_data(as_text=True)
    assert "len zverejnené smeny" in html
    assert "61,00" in html                           # 65 - 4 h konceptu

    # Neplatné parametre nespadnú (použije sa aktuálny mesiac).
    assert admin.get("/reports/hours?year=abc&month=99").status_code == 200
    assert admin.get("/reports/hours?year=1800&month=0").status_code == 200

    assert "nie sú žiadne smeny" in admin.get("/reports/hours?year=2030&month=1").get_data(as_text=True)


def test_employees_cannot_open_reports(hw):
    employee = login(hw, "emp_a", "emppass12")

    for url in ("/reports/hours", "/reports/hours/export.xlsx", "/reports/hours/export.csv"):
        assert employee.get(url).status_code == 403, url


def test_navigation_link_only_for_managers_and_admins(hw):
    for user, password in (("mgr_a", "managerpass1"), ("admin", "adminpass1")):
        assert 'href="/reports/hours"' in login(hw, user, password).get("/").get_data(as_text=True)

    assert 'href="/reports/hours"' not in login(hw, "emp_a", "emppass12").get("/my-schedule").get_data(as_text=True)


def test_xlsx_export(hw):
    from openpyxl import load_workbook

    admin = login(hw, "admin", "adminpass1")
    response = admin.get("/reports/hours/export.xlsx?year=2026&month=9")

    assert response.status_code == 200
    assert "spreadsheetml" in response.mimetype
    assert "prehlad_hodin_2026_09.xlsx" in response.headers["Content-Disposition"]

    workbook = load_workbook(BytesIO(response.data))
    assert workbook.sheetnames == ["Prehľad", "Smeny", "Poznámky"]

    overview = list(workbook["Prehľad"].iter_rows(values_only=True))
    assert overview[0][:4] == ("Zamestnanec", "Pobočka", "Smien", "Hodiny spolu")

    anna = [row for row in overview if row[0] == "Anna Prvá"][0]
    assert anna[1] == "Košice" and anna[2] == 6 and anna[3] == 44
    assert anna[4] == 24 and anna[5] == 8 and anna[6] == 6 and anna[7] == 8
    assert anna[8] == 2                      # dovolenka, prac. dni

    assert overview[-1][0] == "Spolu" and overview[-1][3] == 65

    details = list(workbook["Smeny"].iter_rows(values_only=True))
    assert len(details) - 1 == 9
    assert any(row[0] == "Anna Prvá" and row[13] == "koncept" for row in details[1:])

    notes = " ".join(str(c) for row in workbook["Poznámky"].iter_rows(values_only=True) for c in row if c)
    assert "dátumu jej začiatku" in notes and "Pracovných dní v mesiaci: 20" in notes


def test_xlsx_export_respects_scope_and_published_filter(hw):
    from openpyxl import load_workbook

    manager = login(hw, "mgr_a", "managerpass1")
    workbook = load_workbook(BytesIO(manager.get("/reports/hours/export.xlsx?year=2026&month=9&only_published=1").data))

    names = [row[0] for row in workbook["Prehľad"].iter_rows(min_row=2, values_only=True)]
    assert "Boris Druhý" not in names and "Anna Prvá" in names

    notes = " ".join(str(c) for row in workbook["Poznámky"].iter_rows(values_only=True) for c in row if c)
    assert "len zverejnené smeny" in notes


def test_csv_export(hw):
    admin = login(hw, "admin", "adminpass1")
    response = admin.get("/reports/hours/export.csv?year=2026&month=9")

    assert response.status_code == 200
    assert response.mimetype == "text/csv"
    assert "prehlad_hodin_2026_09.csv" in response.headers["Content-Disposition"]

    text = response.data.decode("utf-8")
    assert text.startswith("\ufeff")

    rows = list(csv.reader(io.StringIO(text.lstrip("\ufeff")), delimiter=";"))
    assert rows[0][0] == "Zamestnanec" and rows[0][3] == "Hodiny spolu"

    anna = [row for row in rows if row[0] == "Anna Prvá"][0]
    assert anna[2] == "6" and anna[3] == "44,00" and anna[4] == "24,00"
    assert rows[-1][0] == "Spolu" and rows[-1][3] == "65,00"

    manager_text = login(hw, "mgr_a", "managerpass1").get("/reports/hours/export.csv?year=2026&month=9").data.decode("utf-8")
    assert "Boris Druhý" not in manager_text
