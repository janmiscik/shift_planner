"""Testy pobočiek a rolí (admin / manažér / zamestnanec).

Každý test dostane čistú SQLite databázu s dvomi pobočkami:
  - "Košice" (A): oddelenie Výroba, zamestnanec Anna (smena 2026-09-21)
  - "Prešov" (B): oddelenie Výroba, zamestnanec Boris (smena 2026-09-22)
a účtami: admin, manager_a (len pobočka A), manager_ab (A aj B),
emp_a (zamestnanec z A), emp_b (zamestnanec z B).
"""

import re

import pytest


def _csrf(html):
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert match is not None, "CSRF token sa v stránke nenašiel"
    return match.group(1)


@pytest.fixture
def world(tmp_path):
    from app.extensions import db
    from app.web.app import create_app

    flask_app = create_app(database_path=tmp_path / "roles.db")
    flask_app.config.update(TESTING=True, WTF_CSRF_ENABLED=True)

    ids = {}

    with flask_app.app_context():
        db.create_all()

        from app.services.auth_service import (
            create_admin,
            create_employee_user,
            create_manager,
        )
        from app.services.branch_service import add_branch
        from app.services.department_service import add_department
        from app.services.employee_department_service import (
            add_employee_department,
        )
        from app.services.employee_service import add_employee
        from app.services.shift_service import add_shift

        ids["a"] = add_branch("Košice")
        ids["b"] = add_branch("Prešov")

        ids["dep_a"] = add_department("Výroba", branch_id=ids["a"])
        ids["dep_b"] = add_department("Výroba", branch_id=ids["b"])

        ids["anna"] = add_employee("Anna", "Prvá", "Operátor", 40, ids["a"])
        ids["boris"] = add_employee("Boris", "Druhý", "Operátor", 40, ids["b"])

        add_employee_department(ids["anna"], ids["dep_a"], 20)
        add_employee_department(ids["boris"], ids["dep_b"], 20)

        # Zverejnené smeny (zamestnanec ich vidí) - nové smeny sú koncept.
        ids["shift_a"] = add_shift(
            ids["anna"], "2026-09-21", "06:00", "14:00", "Ranná",
            department_id=ids["dep_a"], published=True,
        )
        ids["shift_b"] = add_shift(
            ids["boris"], "2026-09-22", "06:00", "14:00", "Ranná",
            department_id=ids["dep_b"], published=True,
        )

        create_admin("admin", "adminpass1")
        create_manager("manager_a", "managerpass1", [ids["a"]])
        create_manager("manager_ab", "managerpass1", [ids["a"], ids["b"]])
        create_employee_user(ids["anna"], "emp_a", "emppass12")
        create_employee_user(ids["boris"], "emp_b", "emppass12")

    flask_app.test_ids = ids

    return flask_app


def login(app, username, password):
    client = app.test_client()
    token = _csrf(client.get("/login").get_data(as_text=True))

    response = client.post(
        "/login",
        data={"csrf_token": token, "username": username, "password": password},
        follow_redirects=True,
    )
    assert response.status_code == 200

    return client


def post(client, page_url, url=None, data=None):
    """POST s CSRF tokenom zo stránky ``page_url`` (bez sledovania
    presmerovania, aby sa dal skontrolovať aj stavový kód)."""

    token = _csrf(client.get(page_url).get_data(as_text=True))
    payload = {"csrf_token": token}
    payload.update(data or {})

    return client.post(url or page_url, data=payload)


# ---------------------------------------------------------------------
# Pobočky (admin)
# ---------------------------------------------------------------------

def test_admin_can_add_branch_and_duplicate_is_rejected(world):
    admin = login(world, "admin", "adminpass1")

    response = post(admin, "/branches/add", data={"name": "Nitra"})
    assert response.status_code == 302

    assert "Nitra" in admin.get("/branches").get_data(as_text=True)

    response = admin.post(
        "/branches/add",
        data={
            "csrf_token": _csrf(admin.get("/branches/add").get_data(as_text=True)),
            "name": "nitra",
        },
    )
    assert "už existuje" in response.get_data(as_text=True)


def test_last_active_branch_cannot_be_deactivated(world):
    from app.extensions import db
    from app.orm_models import Branch

    admin = login(world, "admin", "adminpass1")
    ids = world.test_ids

    # Deaktivujeme pobočku B - ostáva ešte A, takže to prejde.
    post(admin, "/branches", url=f"/branches/toggle/{ids['b']}")

    with world.app_context():
        assert db.session.get(Branch, ids["b"]).active == 0

    # Pobočku A (poslednú aktívnu) deaktivovať nejde.
    post(admin, "/branches", url=f"/branches/toggle/{ids['a']}")

    with world.app_context():
        assert db.session.get(Branch, ids["a"]).active == 1


def test_same_department_name_allowed_in_different_branches(world):
    from app.services.department_service import add_department

    ids = world.test_ids

    with world.app_context():
        # "Výroba" už existuje v A aj B (seed), druhýkrát v A nesmie.
        assert add_department("výroba", branch_id=ids["a"]) is None
        assert add_department("Sklad", branch_id=ids["a"]) is not None
        assert add_department("Sklad", branch_id=ids["b"]) is not None


def test_new_employee_goes_to_selected_branch(world):
    from app.services.employee_service import get_employees

    admin = login(world, "admin", "adminpass1")
    ids = world.test_ids

    response = post(
        admin,
        "/employees/add",
        data={
            "branch_id": str(ids["b"]),
            "first_name": "Cyril",
            "last_name": "Tretí",
            "position": "Skladník",
            "weekly_hours": "40",
            "username": "cyril",
            "password": "cyrilpass1",
        },
    )
    assert response.status_code == 302

    with world.app_context():
        cyril = [e for e in get_employees() if e[1] == "Cyril"][0]
        assert cyril[6] == ids["b"]


# ---------------------------------------------------------------------
# Manažér - rozsah pobočiek
# ---------------------------------------------------------------------

def test_manager_sees_only_own_branch(world):
    manager = login(world, "manager_a", "managerpass1")

    shifts_html = manager.get("/shifts").get_data(as_text=True)
    assert "2026-09-21" in shifts_html
    assert "2026-09-22" not in shifts_html

    employees_html = manager.get("/employees").get_data(as_text=True)
    assert "Anna" in employees_html
    assert "Boris" not in employees_html

    absences_html = manager.get("/absences").get_data(as_text=True)
    assert "Boris" not in absences_html


def test_manager_with_two_branches_sees_both_and_can_filter(world):
    manager = login(world, "manager_ab", "managerpass1")
    ids = world.test_ids

    html = manager.get("/shifts").get_data(as_text=True)
    assert "2026-09-21" in html and "2026-09-22" in html

    html = manager.get(f"/shifts?branch_id={ids['b']}").get_data(as_text=True)
    assert "2026-09-22" in html
    assert "2026-09-21" not in html


def test_manager_cannot_select_foreign_branch_via_query_param(world):
    manager = login(world, "manager_a", "managerpass1")
    ids = world.test_ids

    # Cudzia pobočka sa ignoruje - vráti sa len vlastný rozsah.
    html = manager.get(f"/shifts?branch_id={ids['b']}").get_data(as_text=True)
    assert "2026-09-22" not in html
    assert "2026-09-21" in html

    events = manager.get(
        f"/api/shifts?start=2026-09-01&end=2026-09-30&branch_id={ids['b']}"
    ).get_json()
    assert [e["extendedProps"]["employeeName"] for e in events] == [
        "Anna Prvá"
    ]


def test_manager_cannot_create_shift_for_foreign_branch_employee(world):
    from app.services.shift_service import get_shifts

    manager = login(world, "manager_a", "managerpass1")
    ids = world.test_ids


    response = post(
        manager,
        "/shifts/add",
        data={
            "employee_id": str(ids["boris"]),
            "department_id": "",
            "shift_date": "2026-09-23",
            "start_time": "06:00",
            "end_time": "14:00",
            "shift_type": "Ranná",
        },
    )

    # Formulár sa vráti s chybou (zamestnanec nie je v ponuke), nie 302.
    assert response.status_code == 200

    with world.app_context():
        dates = [s[3] for s in get_shifts()]
        assert "2026-09-23" not in dates


def test_manager_can_create_shift_in_own_branch(world):
    from app.services.shift_service import get_shifts

    manager = login(world, "manager_a", "managerpass1")
    ids = world.test_ids

    response = post(
        manager,
        "/shifts/add",
        data={
            "employee_id": str(ids["anna"]),
            "department_id": str(ids["dep_a"]),
            "shift_date": "2026-09-23",
            "start_time": "06:00",
            "end_time": "14:00",
            "shift_type": "Ranná",
        },
    )
    assert response.status_code == 302

    with world.app_context():
        assert "2026-09-23" in [s[3] for s in get_shifts()]


def test_manager_cannot_edit_delete_or_move_foreign_shift(world):
    from app.services.shift_service import get_shifts

    manager = login(world, "manager_a", "managerpass1")
    ids = world.test_ids
    foreign = ids["shift_b"]

    assert manager.get(f"/shifts/edit/{foreign}").status_code == 404

    response = post(manager, "/shifts", url=f"/shifts/delete/{foreign}")
    assert response.status_code == 404

    with world.app_context():
        assert foreign in [s[0] for s in get_shifts()]

    token = _csrf(manager.get("/calendar").get_data(as_text=True))
    response = manager.post(
        f"/api/shifts/{foreign}/move",
        json={
            "shift_date": "2026-09-25",
            "start_time": "06:00",
            "end_time": "14:00",
        },
        headers={"X-CSRFToken": token},
    )
    assert response.status_code == 404


def test_manager_cannot_touch_foreign_absence_or_employee_data(world):
    from app.services.absence_service import add_absence, get_absences

    ids = world.test_ids

    with world.app_context():
        foreign_absence = add_absence(
            ids["boris"], "Dovolenka", "2026-10-01", "2026-10-02"
        )

    manager = login(world, "manager_a", "managerpass1")

    assert manager.get(f"/absences/edit/{foreign_absence}").status_code == 404
    assert (
        post(manager, "/absences", url=f"/absences/delete/{foreign_absence}")
        .status_code
        == 404
    )

    with world.app_context():
        assert foreign_absence in [a[0] for a in get_absences()]

    assert (
        manager.get(f"/api/employees/{ids['boris']}/departments").status_code
        == 404
    )
    assert (
        manager.get(f"/employees/{ids['boris']}/export.ics").status_code == 404
    )
    assert (
        manager.get(f"/employees/{ids['anna']}/export.ics").status_code == 200
    )


def test_manager_exports_contain_only_own_branch(world):
    from openpyxl import load_workbook
    from io import BytesIO

    manager = login(world, "manager_a", "managerpass1")

    response = manager.get("/shifts/export.xlsx")
    assert response.status_code == 200

    sheet = load_workbook(BytesIO(response.data)).active
    text = " ".join(
        str(cell) for row in sheet.iter_rows(values_only=True) for cell in row
    )
    assert "Anna" in text
    assert "Boris" not in text

    assert manager.get("/shifts/export.pdf").status_code == 200
    assert "Boris" not in manager.get("/shifts/export.ics").get_data(
        as_text=True
    )
    assert "Boris" not in manager.get(
        "/calendar/export.ics?year=2026&month=9"
    ).get_data(as_text=True)


def test_shift_with_department_from_other_branch_is_rejected(world):
    from app.services.shift_service import validate_shift

    ids = world.test_ids

    with world.app_context():
        errors = validate_shift(
            ids["anna"],
            "2026-09-24",
            "06:00",
            "14:00",
            department_id=ids["dep_b"],
        )

    assert any("inej pobočky" in error for error in errors)


def test_employee_department_choices_limited_to_employee_branch(world):
    admin = login(world, "admin", "adminpass1")
    ids = world.test_ids

    # Oddelenie z cudzej pobočky nie je v ponuke -> priradenie zlyhá.
    response = post(
        admin,
        f"/employees/{ids['anna']}/departments/add",
        data={"department_id": str(ids["dep_b"]), "weekly_hours": "5"},
    )
    assert response.status_code == 200

    with world.app_context():
        from app.services.employee_department_service import (
            get_employee_department_ids,
        )

        assert get_employee_department_ids(ids["anna"]) == {ids["dep_a"]}


# ---------------------------------------------------------------------
# Manažér - zakázané routy
# ---------------------------------------------------------------------

@pytest.mark.parametrize(
    "url",
    [
        "/branches",
        "/branches/add",
        "/users",
        "/users/add",
        "/admin/backup",
        "/employees/add",
        "/departments/add",
        "/my-schedule",
    ],
)
def test_manager_is_forbidden_from_admin_routes(world, url):
    manager = login(world, "manager_ab", "managerpass1")

    assert manager.get(url).status_code in (302, 403)

    if url != "/my-schedule":
        assert manager.get(url).status_code == 403


def test_manager_cannot_modify_employees_or_departments(world):
    manager = login(world, "manager_a", "managerpass1")
    ids = world.test_ids

    token = _csrf(manager.get("/shifts").get_data(as_text=True))

    for url in (
        f"/employees/toggle/{ids['anna']}",
        f"/departments/toggle/{ids['dep_a']}",
        f"/employees/edit/{ids['anna']}",
        f"/users/delete/1",
    ):
        response = manager.post(url, data={"csrf_token": token})
        assert response.status_code == 403, url


def test_manager_pages_hide_admin_actions(world):
    manager = login(world, "manager_a", "managerpass1")

    html = manager.get("/employees").get_data(as_text=True)
    assert "Pridať zamestnanca" not in html
    assert "/employees/edit/" not in html

    html = manager.get("/").get_data(as_text=True)
    assert "/branches" not in html
    assert "/users" not in html
    assert "/admin/backup" not in html


# ---------------------------------------------------------------------
# Zamestnanec - čítanie a export
# ---------------------------------------------------------------------

def test_employee_can_view_and_export_own_schedule(world):
    employee = login(world, "emp_a", "emppass12")

    html = employee.get("/my-schedule").get_data(as_text=True)
    assert "2026-09-21" in html
    assert "2026-09-22" not in html

    for url, mimetype in (
        ("/my-schedule/export.ics", "text/calendar"),
        ("/my-schedule/export.xlsx", "spreadsheetml"),
        ("/my-schedule/export.pdf", "application/pdf"),
    ):
        response = employee.get(url)
        assert response.status_code == 200, url
        assert mimetype in response.mimetype, url


def test_employee_branch_calendar_is_read_only_and_branch_scoped(
    world, monkeypatch
):
    # Voliteľná funkcia - predvolene je vypnutá.
    monkeypatch.setenv("EMPLOYEES_CAN_VIEW_BRANCH_CALENDAR", "1")

    employee = login(world, "emp_a", "emppass12")

    html = employee.get("/calendar").get_data(as_text=True)
    assert "const readOnly = true;" in html
    assert "Pridať smenu" not in html

    events = employee.get("/api/shifts?start=2026-09-01&end=2026-09-30").get_json()
    assert [e["extendedProps"]["employeeName"] for e in events] == [
        "Anna Prvá"
    ]
    assert all(e["extendedProps"]["weeklyHours"] is None for e in events)

    # Ani cez ?branch_id= sa nedostane do cudzej pobočky.
    ids = world.test_ids
    events = employee.get(
        f"/api/shifts?start=2026-09-01&end=2026-09-30&branch_id={ids['b']}"
    ).get_json()
    assert [e["extendedProps"]["employeeName"] for e in events] == [
        "Anna Prvá"
    ]

    assert employee.get("/calendar/export.xlsx?year=2026&month=9").status_code == 200


def test_employee_cannot_use_planning_routes(world):
    employee = login(world, "emp_a", "emppass12")
    ids = world.test_ids

    for url in ("/shifts", "/shifts/add", "/employees", "/absences", "/branches", "/users", "/admin/backup", "/calendar", "/api/shifts?start=2026-09-01&end=2026-09-30"):
        assert employee.get(url).status_code == 403, url

    token = _csrf(employee.get("/my-schedule").get_data(as_text=True))
    response = employee.post(
        f"/api/shifts/{ids['shift_a']}/move",
        json={"shift_date": "2026-09-25"},
        headers={"X-CSRFToken": token},
    )
    assert response.status_code == 403


def test_employee_sees_only_own_shifts_by_default(world):
    """Predvolene zamestnanec nevidí smeny ostatných - ani v
    kalendári, ani v API."""

    employee = login(world, "emp_a", "emppass12")

    assert employee.get("/calendar").status_code == 403
    assert employee.get("/api/shifts?start=2026-09-01&end=2026-09-30").status_code == 403
    assert employee.get("/calendar/export.xlsx?year=2026&month=9").status_code == 403

    html = employee.get("/my-schedule").get_data(as_text=True)
    assert "2026-09-21" in html
    assert "2026-09-22" not in html
    assert "Boris" not in html


# ---------------------------------------------------------------------
# Správa účtov (admin)
# ---------------------------------------------------------------------

def test_admin_creates_manager_with_branches(world):
    from app.services.auth_service import get_user_by_username

    admin = login(world, "admin", "adminpass1")
    ids = world.test_ids

    response = admin.post(
        "/users/add",
        data={
            "csrf_token": _csrf(admin.get("/users/add").get_data(as_text=True)),
            "username": "novy_manager",
            "password": "novypass12",
            "role": "manager",
            "branch_ids": [str(ids["b"])],
        },
    )
    assert response.status_code == 302

    with world.app_context():
        user = get_user_by_username("novy_manager")
        assert user.role == "manager"
        assert [b.id for b in user.branches] == [ids["b"]]

    # Nový manažér vidí len pobočku B.
    manager = login(world, "novy_manager", "novypass12")
    html = manager.get("/shifts").get_data(as_text=True)
    assert "2026-09-22" in html
    assert "2026-09-21" not in html


def test_manager_without_branch_is_rejected(world):
    from app.services.auth_service import get_user_by_username

    admin = login(world, "admin", "adminpass1")

    response = admin.post(
        "/users/add",
        data={
            "csrf_token": _csrf(admin.get("/users/add").get_data(as_text=True)),
            "username": "bez_pobocky",
            "password": "novypass12",
            "role": "manager",
        },
    )
    assert "aspoň jednu pobočku" in response.get_data(as_text=True)

    with world.app_context():
        assert get_user_by_username("bez_pobocky") is None


def test_admin_can_change_manager_branches(world):
    from app.services.auth_service import get_user_by_username

    admin = login(world, "admin", "adminpass1")
    ids = world.test_ids

    with world.app_context():
        user_id = get_user_by_username("manager_a").id

    response = admin.post(
        f"/users/edit/{user_id}",
        data={
            "csrf_token": _csrf(
                admin.get(f"/users/edit/{user_id}").get_data(as_text=True)
            ),
            "username": "manager_a",
            "password": "",
            "branch_ids": [str(ids["b"])],
        },
    )
    assert response.status_code == 302

    with world.app_context():
        user = get_user_by_username("manager_a")
        assert [b.id for b in user.branches] == [ids["b"]]


def test_cannot_delete_self_or_last_admin(world):
    from app.services.auth_service import count_admins, get_user_by_username

    admin = login(world, "admin", "adminpass1")

    with world.app_context():
        admin_id = get_user_by_username("admin").id

    response = post(admin, "/users", url=f"/users/delete/{admin_id}")
    assert response.status_code == 302

    with world.app_context():
        assert count_admins() == 1
        assert get_user_by_username("admin") is not None


def test_admin_can_delete_manager(world):
    from app.services.auth_service import get_user_by_username

    admin = login(world, "admin", "adminpass1")

    with world.app_context():
        user_id = get_user_by_username("manager_ab").id

    post(admin, "/users", url=f"/users/delete/{user_id}")

    with world.app_context():
        assert get_user_by_username("manager_ab") is None


# ---------------------------------------------------------------------
# Funkcia manažéra na zamestnaneckom účte (jeden účet)
# ---------------------------------------------------------------------

def _set_manager_function(admin, employee_id, enabled, branch_ids):
    data = {"branch_ids": [str(b) for b in branch_ids]}

    if enabled:
        data["manager_enabled"] = "1"

    token = _csrf(
        admin.get(f"/employees/edit/{employee_id}").get_data(as_text=True)
    )
    data["csrf_token"] = token

    return admin.post(f"/employees/{employee_id}/manager", data=data)


def test_admin_gives_employee_account_manager_function(world):
    from app.services.auth_service import get_user_by_username

    admin = login(world, "admin", "adminpass1")
    ids = world.test_ids

    # Pred prideleniem pracovník plánovať nesmie.
    worker = login(world, "emp_a", "emppass12")
    assert worker.get("/shifts").status_code == 403

    response = _set_manager_function(admin, ids["anna"], True, [ids["a"]])
    assert response.status_code == 302

    with world.app_context():
        user = get_user_by_username("emp_a")
        assert user.role == "manager"
        assert user.employee_id == ids["anna"]
        assert [b.id for b in user.branches] == [ids["a"]]

    # Ten istý účet: plánuje v pobočke A ...
    worker = login(world, "emp_a", "emppass12")
    html = worker.get("/shifts").get_data(as_text=True)
    assert "2026-09-21" in html
    assert "2026-09-22" not in html

    response = post(
        worker,
        "/shifts/add",
        data={
            "employee_id": str(ids["anna"]),
            "department_id": str(ids["dep_a"]),
            "shift_date": "2026-09-23",
            "start_time": "06:00",
            "end_time": "14:00",
            "shift_type": "Ranná",
        },
    )
    assert response.status_code == 302

    # ... do pobočky B nevidí ani nezasiahne ...
    assert worker.get(f"/shifts/edit/{ids['shift_b']}").status_code == 404

    # ... a naďalej vidí vlastný rozpis (len svoje smeny).
    html = worker.get("/my-schedule").get_data(as_text=True)
    assert "2026-09-21" in html
    assert "2026-09-22" not in html
    assert worker.get("/my-schedule/export.xlsx").status_code == 200
    assert 'href="/my-schedule"' in worker.get("/").get_data(as_text=True)

    # Admin veci ostávajú zakázané.
    for url in ("/users", "/branches", "/admin/backup", "/employees/add"):
        assert worker.get(url).status_code == 403, url


def test_admin_can_take_manager_function_back(world):
    admin = login(world, "admin", "adminpass1")
    ids = world.test_ids

    _set_manager_function(admin, ids["anna"], True, [ids["a"]])
    _set_manager_function(admin, ids["anna"], False, [])

    worker = login(world, "emp_a", "emppass12")

    assert worker.get("/shifts").status_code == 403
    assert worker.get("/my-schedule").status_code == 200

    with world.app_context():
        from app.services.auth_service import get_user_by_username

        user = get_user_by_username("emp_a")
        assert user.role == "employee"
        assert user.branches == []


def test_manager_function_requires_branch_and_account(world):
    from app.services.employee_service import add_employee
    from app.services.auth_service import get_user_by_username

    admin = login(world, "admin", "adminpass1")
    ids = world.test_ids

    # Bez pobočky sa funkcia nepridelí.
    response = _set_manager_function(admin, ids["anna"], True, [])
    assert response.status_code == 302

    with world.app_context():
        assert get_user_by_username("emp_a").role == "employee"

    # Zamestnanec bez účtu: najprv treba vytvoriť účet.
    with world.app_context():
        no_account = add_employee("Dana", "Bez", "Op", 40, ids["a"])

    response = admin.post(
        f"/employees/{no_account}/manager",
        data={
            "csrf_token": _csrf(
                admin.get(f"/employees/edit/{ids['anna']}").get_data(as_text=True)
            ),
            "manager_enabled": "1",
            "branch_ids": [str(ids["a"])],
        },
    )
    assert response.status_code == 302


def test_only_admin_can_assign_manager_function(world):
    manager = login(world, "manager_a", "managerpass1")
    ids = world.test_ids

    token = _csrf(manager.get("/shifts").get_data(as_text=True))
    response = manager.post(
        f"/employees/{ids['anna']}/manager",
        data={
            "csrf_token": token,
            "manager_enabled": "1",
            "branch_ids": [str(ids["a"])],
        },
    )
    assert response.status_code == 403

    worker = login(world, "emp_b", "emppass12")
    token = _csrf(worker.get("/my-schedule").get_data(as_text=True))
    response = worker.post(
        f"/employees/{ids['boris']}/manager",
        data={
            "csrf_token": token,
            "manager_enabled": "1",
            "branch_ids": [str(ids["b"])],
        },
    )
    assert response.status_code == 403


# ---------------------------------------------------------------------
# Zmena hesla
# ---------------------------------------------------------------------

def test_user_can_change_own_password(world):
    manager = login(world, "manager_a", "managerpass1")

    response = post(
        manager,
        "/account/password",
        data={
            "current_password": "managerpass1",
            "new_password": "novyheslo1",
            "confirm_password": "novyheslo1",
        },
    )
    assert response.status_code == 302

    login(world, "manager_a", "novyheslo1").get("/shifts")

    client = world.test_client()
    token = _csrf(client.get("/login").get_data(as_text=True))
    response = client.post(
        "/login",
        data={
            "csrf_token": token,
            "username": "manager_a",
            "password": "managerpass1",
        },
        follow_redirects=True,
    )
    assert "Nesprávne" in response.get_data(as_text=True)


def test_change_password_rejects_wrong_current_password(world):
    employee = login(world, "emp_a", "emppass12")

    response = post(
        employee,
        "/account/password",
        data={
            "current_password": "zle",
            "new_password": "novyheslo1",
            "confirm_password": "novyheslo1",
        },
    )
    assert "nie je správne" in response.get_data(as_text=True)


# ---------------------------------------------------------------------
# Kontrola, že každá routa má rozhodnuté oprávnenie
# ---------------------------------------------------------------------

def test_every_endpoint_has_an_explicit_permission_decision(world):
    """Ak pridáš novú routu, tento test zlyhá, kým nerozhodneš, či ju
    smie používať manažér/zamestnanec (zoznamy v app/web/app.py),
    alebo je len pre admina (pridaj ju do ``ADMIN_ONLY`` nižšie)."""

    from app.web import app as app_module

    ADMIN_ONLY = {
        "main.admin_backup",
        "main.add_branch_page",
        "main.add_department_page",
        "main.add_employee_department_page",
        "main.add_employee_page",
        "main.add_user_page",
        "main.branches_page",
        "main.create_employee_credentials_page",
        "main.delete_employee_credentials_page",
        "main.delete_employee_department_page",
        "main.delete_user_page",
        "main.edit_branch_page",
        "main.edit_department_page",
        "main.edit_employee_department_page",
        "main.edit_employee_page",
        "main.edit_user_page",
        "main.toggle_branch_page",
        "main.toggle_department_page",
        "main.toggle_employee_page",
        "main.update_employee_credentials_page",
        "main.update_employee_manager_page",
        "main.users_page",
    }

    all_endpoints = {
        rule.endpoint for rule in world.url_map.iter_rules()
    } - app_module._PUBLIC_ENDPOINTS

    decided = (
        app_module._MANAGER_ALLOWED_ENDPOINTS
        | app_module._EMPLOYEE_ALLOWED_ENDPOINTS
        | app_module._EMPLOYEE_BRANCH_VIEW_ENDPOINTS
        | ADMIN_ONLY
    )

    # Oprávnenia nesmú odkazovať na neexistujúce routy (preklep).
    assert (
        app_module._MANAGER_ALLOWED_ENDPOINTS
        | app_module._EMPLOYEE_ALLOWED_ENDPOINTS
        | app_module._EMPLOYEE_BRANCH_VIEW_ENDPOINTS
        | ADMIN_ONLY
    ) <= all_endpoints

    assert all_endpoints - decided == set()

    # Admin-only routy nesmie vidieť manažér ani zamestnanec.
    assert not (ADMIN_ONLY & app_module._MANAGER_ALLOWED_ENDPOINTS)
    assert not (ADMIN_ONLY & app_module._EMPLOYEE_ALLOWED_ENDPOINTS)
    assert not (ADMIN_ONLY & app_module._EMPLOYEE_BRANCH_VIEW_ENDPOINTS)


# ---------------------------------------------------------------------
# Migrácia existujúcej databázy
# ---------------------------------------------------------------------

def test_migration_moves_existing_data_to_default_branch(tmp_path):
    """Databáza vo verzii pred pobočkami (so smenami, absenciami a
    manažérom) sa po migrácii zachová: všetko je v "Hlavnej
    pobočke" a pôvodný manažér sa stane adminom."""

    import sqlalchemy as sa
    from flask_migrate import upgrade

    from app.extensions import db
    from app.web.app import create_app

    flask_app = create_app(database_path=tmp_path / "migrate.db")

    with flask_app.app_context():
        upgrade(revision="aeca62578799")

        connection = db.session.connection()
        connection.execute(sa.text(
            "INSERT INTO departments (name, active, min_staff) "
            "VALUES ('Vyroba', 1, 2)"
        ))
        connection.execute(sa.text(
            "INSERT INTO employees (first_name, last_name, position, "
            "employment_type, weekly_hours, active) "
            "VALUES ('Jan', 'Novak', 'Op', 'custom', 40, 1)"
        ))
        connection.execute(sa.text(
            "INSERT INTO employee_departments "
            "(employee_id, department_id, weekly_hours) VALUES (1, 1, 20)"
        ))
        connection.execute(sa.text(
            "INSERT INTO shifts (employee_id, shift_date, start_time, "
            "end_time, shift_type, department_id) "
            "VALUES (1, '2026-09-21', '06:00', '14:00', 'Ranná', 1)"
        ))
        connection.execute(sa.text(
            "INSERT INTO absences (employee_id, absence_type, start_date, "
            "end_date) VALUES (1, 'PN', '2026-10-01', '2026-10-02')"
        ))
        connection.execute(sa.text(
            "INSERT INTO users (username, password_hash, role, employee_id) "
            "VALUES ('boss', 'x', 'manager', NULL), ('jan', 'x', 'employee', 1)"
        ))
        db.session.commit()

        upgrade()

        connection = db.session.connection()

        branches = connection.execute(
            sa.text("SELECT id, name FROM branches")
        ).all()
        assert branches == [(1, "Hlavná pobočka")]

        assert connection.execute(
            sa.text("SELECT branch_id FROM departments")
        ).scalars().all() == [1]
        assert connection.execute(
            sa.text("SELECT branch_id FROM employees")
        ).scalars().all() == [1]

        for table, expected in (
            ("shifts", 1),
            ("absences", 1),
            ("employee_departments", 1),
        ):
            assert (
                connection.execute(sa.text(f"SELECT COUNT(*) FROM {table}")).scalar()
                == expected
            ), table

        roles = dict(
            connection.execute(sa.text("SELECT username, role FROM users")).all()
        )
        assert roles == {"boss": "admin", "jan": "employee"}

        if connection.dialect.name == "sqlite":
            assert (
                connection.execute(sa.text("PRAGMA foreign_key_check")).all()
                == []
            )

        # Po migrácii sa názov oddelenia môže opakovať v inej pobočke,
        # ale nie v tej istej.
        connection.execute(sa.text(
            "INSERT INTO branches (name, active) VALUES ('Nitra', 1)"
        ))
        connection.execute(sa.text(
            "INSERT INTO departments (name, active, branch_id) "
            "VALUES ('Vyroba', 1, 2)"
        ))

        with pytest.raises(sa.exc.IntegrityError):
            connection.execute(sa.text(
                "INSERT INTO departments (name, active, branch_id) "
                "VALUES ('Vyroba', 1, 2)"
            ))

        db.session.rollback()


# ---------------------------------------------------------------------
# Všetky stránky sa vykreslia (bez chýb v šablónach) pre každú rolu
# ---------------------------------------------------------------------

def test_all_pages_render_for_each_role(world):
    ids = world.test_ids

    admin_pages = [
        "/", "/employees", "/employees/add",
        f"/employees/edit/{ids['anna']}",
        f"/employees/{ids['anna']}/departments/add",
        "/departments", "/departments/add", f"/departments/edit/{ids['dep_a']}",
        "/shifts", "/shifts/add", f"/shifts/edit/{ids['shift_a']}",
        "/absences", "/absences/add", "/calendar",
        "/branches", "/branches/add", f"/branches/edit/{ids['a']}",
        "/users", "/users/add", "/account/password",
        f"/shifts?branch_id={ids['a']}", f"/calendar?branch_id={ids['b']}",
    ]
    admin = login(world, "admin", "adminpass1")

    for url in admin_pages:
        assert admin.get(url).status_code == 200, url

    manager_pages = [
        "/", "/employees", "/departments", "/shifts", "/shifts/add",
        f"/shifts/edit/{ids['shift_a']}", "/absences", "/absences/add",
        "/calendar", "/account/password",
    ]
    manager = login(world, "manager_ab", "managerpass1")

    for url in manager_pages:
        assert manager.get(url).status_code == 200, url

    employee = login(world, "emp_a", "emppass12")

    for url in ("/my-schedule", "/account/password"):
        assert employee.get(url).status_code == 200, url

    # Admin s viacerými pobočkami vidí v zozname zamestnancov aj pobočku.
    html = admin.get("/employees").get_data(as_text=True)
    assert "Košice" in html and "Prešov" in html
