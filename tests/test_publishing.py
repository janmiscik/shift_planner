"""Testy zverejňovania rozpisu (koncept vs. zverejnené).

Scenár: pobočka A "Košice" (Anna, manažér mgr_a) a pobočka B "Prešov"
(Boris). Smeny sú vytvorené bežným spôsobom, teda ako KONCEPT:
  - Anna: po 21.9., ut 22.9. (týždeň 21.-27.9.)
  - Boris: po 21.9.
  - Anna: po 28.9. - už zverejnená (iný týždeň)
"""

import re
from io import BytesIO

import pytest


def _csrf(html):
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert match is not None
    return match.group(1)


@pytest.fixture
def pub_world(tmp_path, monkeypatch):
    monkeypatch.delenv("AUTO_PUBLISH_SHIFTS", raising=False)

    from app.extensions import db
    from app.web.app import create_app

    app = create_app(database_path=tmp_path / "pub.db")
    app.config.update(TESTING=True, WTF_CSRF_ENABLED=True)

    ids = {}

    with app.app_context():
        db.create_all()

        from app.services.auth_service import (
            create_admin,
            create_employee_user,
            create_manager,
        )
        from app.services.branch_service import add_branch
        from app.services.employee_service import add_employee
        from app.services.shift_service import add_shift

        ids["a"] = add_branch("Košice")
        ids["b"] = add_branch("Prešov")
        ids["anna"] = add_employee("Anna", "Prvá", "Op", 40, ids["a"])
        ids["boris"] = add_employee("Boris", "Druhý", "Op", 40, ids["b"])

        ids["s_mon"] = add_shift(ids["anna"], "2026-09-21", "06:00", "14:00", "Ranná")
        ids["s_tue"] = add_shift(ids["anna"], "2026-09-22", "06:00", "14:00", "Ranná")
        ids["s_boris"] = add_shift(ids["boris"], "2026-09-21", "06:00", "14:00", "Ranná")
        ids["s_next"] = add_shift(
            ids["anna"], "2026-09-28", "06:00", "14:00", "Ranná", published=True
        )

        create_admin("admin", "adminpass1")
        create_manager("mgr_a", "managerpass1", [ids["a"]])
        create_employee_user(ids["anna"], "emp_a", "emppass12")
        create_employee_user(ids["boris"], "emp_b", "emppass12")

    app.test_ids = ids

    return app


def login(app, username, password):
    client = app.test_client()
    token = _csrf(client.get("/login").get_data(as_text=True))
    client.post(
        "/login",
        data={"csrf_token": token, "username": username, "password": password},
    )
    return client


def publish_action(client, app, action, branch="a", week="2026-09-21", **extra):
    token = _csrf(client.get("/shifts/publish").get_data(as_text=True))

    return client.post(
        "/shifts/publish",
        data={
            "csrf_token": token,
            "action": action,
            "branch_id": str(app.test_ids[branch]),
            "week_date": week,
            **extra,
        },
    )


def published_flags(app):
    from app.extensions import db
    from app.orm_models import Shift

    with app.app_context():
        return {s.id: s.published for s in db.session.query(Shift).all()}


# ---------------------------------------------------------------------
# Predvolene koncept
# ---------------------------------------------------------------------

def test_new_shifts_are_draft_by_default(pub_world):
    flags = published_flags(pub_world)
    ids = pub_world.test_ids

    assert flags[ids["s_mon"]] == 0
    assert flags[ids["s_tue"]] == 0
    assert flags[ids["s_next"]] == 1


def test_shift_added_in_the_app_is_a_draft(pub_world):
    manager = login(pub_world, "mgr_a", "managerpass1")
    ids = pub_world.test_ids

    token = _csrf(manager.get("/shifts/add").get_data(as_text=True))
    response = manager.post(
        "/shifts/add",
        data={
            "csrf_token": token,
            "employee_id": str(ids["anna"]),
            "department_id": "",
            "shift_date": "2026-09-23",
            "start_time": "06:00",
            "end_time": "14:00",
            "shift_type": "Ranná",
        },
    )
    assert response.status_code == 302

    flags = published_flags(pub_world)
    new_id = max(flags)

    assert flags[new_id] == 0


def test_auto_publish_setting_publishes_new_shifts_immediately(
    pub_world, monkeypatch
):
    from app.services.shift_service import add_shift

    monkeypatch.setenv("AUTO_PUBLISH_SHIFTS", "1")
    ids = pub_world.test_ids

    with pub_world.app_context():
        shift_id = add_shift(ids["anna"], "2026-10-05", "06:00", "14:00", "Ranná")

    assert published_flags(pub_world)[shift_id] == 1


# ---------------------------------------------------------------------
# Zamestnanec vidí len zverejnené
# ---------------------------------------------------------------------

def test_employee_does_not_see_drafts_anywhere(pub_world):
    from openpyxl import load_workbook

    employee = login(pub_world, "emp_a", "emppass12")

    html = employee.get("/my-schedule").get_data(as_text=True)
    assert "2026-09-21" not in html and "2026-09-22" not in html
    assert "2026-09-28" in html   # zverejnená smena je vidieť
    assert "zverejnené manažérom" in html

    ics = employee.get("/my-schedule/export.ics").get_data(as_text=True)
    assert ics.count("BEGIN:VEVENT") == 1

    sheet = load_workbook(BytesIO(employee.get("/my-schedule/export.xlsx").data)).active
    text = " ".join(
        str(cell) for row in sheet.iter_rows(values_only=True) for cell in row
    )
    assert "2026-09-21" not in text and "2026-09-28" in text

    assert employee.get("/my-schedule/export.pdf").status_code == 200


def test_employee_branch_calendar_hides_drafts(pub_world, monkeypatch):
    monkeypatch.setenv("EMPLOYEES_CAN_VIEW_BRANCH_CALENDAR", "1")

    employee = login(pub_world, "emp_a", "emppass12")

    events = employee.get("/api/shifts?start=2026-09-01&end=2026-10-31").get_json()
    assert [e["start"][:10] for e in events] == ["2026-09-28"]
    assert not any("shift-draft" in e["className"] for e in events)

    ics = employee.get("/calendar/export.ics?year=2026&month=9").get_data(as_text=True)
    assert ics.count("BEGIN:VEVENT") == 1


# ---------------------------------------------------------------------
# Zverejnenie / stiahnutie
# ---------------------------------------------------------------------

def test_publishing_a_week_makes_shifts_visible_to_the_employee(pub_world):
    manager = login(pub_world, "mgr_a", "managerpass1")
    employee = login(pub_world, "emp_a", "emppass12")
    ids = pub_world.test_ids

    response = publish_action(manager, pub_world, "publish")
    assert response.status_code == 302

    flags = published_flags(pub_world)
    assert flags[ids["s_mon"]] == 1 and flags[ids["s_tue"]] == 1

    html = employee.get("/my-schedule").get_data(as_text=True)
    assert "2026-09-21" in html and "2026-09-22" in html


def test_publish_is_limited_to_the_chosen_week_and_branch(pub_world):
    admin = login(pub_world, "admin", "adminpass1")
    ids = pub_world.test_ids

    # Do týždňa 28.9. pridáme koncept - nesmie sa zverejniť s týždňom 21.9.
    from app.services.shift_service import add_shift

    with pub_world.app_context():
        later_draft = add_shift(ids["anna"], "2026-09-29", "06:00", "14:00", "Ranná")

    publish_action(admin, pub_world, "publish", branch="a", week="2026-09-23")

    flags = published_flags(pub_world)
    assert flags[ids["s_mon"]] == 1 and flags[ids["s_tue"]] == 1
    assert flags[ids["s_boris"]] == 0      # iná pobočka
    assert flags[later_draft] == 0         # iný týždeň


def test_unpublish_hides_shifts_from_the_employee_again(pub_world):
    manager = login(pub_world, "mgr_a", "managerpass1")
    employee = login(pub_world, "emp_a", "emppass12")

    publish_action(manager, pub_world, "publish")
    assert "2026-09-21" in employee.get("/my-schedule").get_data(as_text=True)

    publish_action(manager, pub_world, "unpublish")
    assert "2026-09-21" not in employee.get("/my-schedule").get_data(as_text=True)


def test_summary_counts_drafts_and_published(pub_world):
    from datetime import date

    from app.services.publish_service import week_publication_summary

    ids = pub_world.test_ids

    with pub_world.app_context():
        assert week_publication_summary(date(2026, 9, 21), ids["a"]) == {
            "total": 2, "drafts": 2, "published": 0,
        }
        assert week_publication_summary(date(2026, 9, 28), ids["a"]) == {
            "total": 1, "drafts": 0, "published": 1,
        }
        assert week_publication_summary(date(2026, 9, 21), ids["b"])["total"] == 1


def test_publish_page_shows_state_and_actions(pub_world):
    manager = login(pub_world, "mgr_a", "managerpass1")
    ids = pub_world.test_ids

    html = manager.get(
        f"/shifts/publish?branch_id={ids['a']}&week_date=2026-09-23"
    ).get_data(as_text=True)

    assert "koncept:" in html and "Zverejniť 2 smien" in html
    assert "Stiahnuť zo zverejnenia" not in html   # nič zverejnené

    publish_action(manager, pub_world, "publish")

    html = manager.get(
        f"/shifts/publish?branch_id={ids['a']}&week_date=2026-09-23"
    ).get_data(as_text=True)
    assert "Stiahnuť zo zverejnenia (2)" in html
    assert "už zverejnené" in html


# ---------------------------------------------------------------------
# Úpravy po zverejnení
# ---------------------------------------------------------------------

def test_editing_or_moving_a_published_shift_keeps_it_published(pub_world):
    manager = login(pub_world, "mgr_a", "managerpass1")
    ids = pub_world.test_ids

    publish_action(manager, pub_world, "publish")

    token = _csrf(manager.get(f"/shifts/edit/{ids['s_mon']}").get_data(as_text=True))
    response = manager.post(
        f"/shifts/edit/{ids['s_mon']}",
        data={
            "csrf_token": token,
            "employee_id": str(ids["anna"]),
            "department_id": "",
            "shift_date": "2026-09-21",
            "start_time": "07:00",
            "end_time": "15:00",
            "shift_type": "Ranná",
        },
    )
    assert response.status_code == 302
    assert published_flags(pub_world)[ids["s_mon"]] == 1

    token = _csrf(manager.get("/calendar").get_data(as_text=True))
    response = manager.post(
        f"/api/shifts/{ids['s_tue']}/move",
        json={"shift_date": "2026-09-24", "start_time": "06:00", "end_time": "14:00"},
        headers={"X-CSRFToken": token},
    )
    assert response.status_code == 200
    assert published_flags(pub_world)[ids["s_tue"]] == 1


def test_shift_added_to_a_published_week_is_a_draft_until_republished(pub_world):
    manager = login(pub_world, "mgr_a", "managerpass1")
    employee = login(pub_world, "emp_a", "emppass12")
    ids = pub_world.test_ids

    publish_action(manager, pub_world, "publish")

    token = _csrf(manager.get("/shifts/add").get_data(as_text=True))
    manager.post(
        "/shifts/add",
        data={
            "csrf_token": token,
            "employee_id": str(ids["anna"]),
            "department_id": "",
            "shift_date": "2026-09-25",
            "start_time": "06:00",
            "end_time": "14:00",
            "shift_type": "Ranná",
        },
    )

    assert "2026-09-25" not in employee.get("/my-schedule").get_data(as_text=True)

    publish_action(manager, pub_world, "publish")
    assert "2026-09-25" in employee.get("/my-schedule").get_data(as_text=True)


# ---------------------------------------------------------------------
# Manažér vidí stav
# ---------------------------------------------------------------------

def test_manager_sees_status_in_list_and_calendar(pub_world):
    manager = login(pub_world, "mgr_a", "managerpass1")
    ids = pub_world.test_ids

    html = manager.get("/shifts").get_data(as_text=True)
    assert "Koncept" in html and "Zverejnená" in html

    events = manager.get("/api/shifts?start=2026-09-01&end=2026-10-31").get_json()
    by_date = {e["start"][:10]: e for e in events}

    assert by_date["2026-09-21"]["extendedProps"]["draft"] is True
    assert "shift-draft" in by_date["2026-09-21"]["className"]
    assert by_date["2026-09-28"]["extendedProps"]["draft"] is False
    assert "shift-draft" not in by_date["2026-09-28"]["className"]

    assert "Koncept (zamestnanec ho zatiaľ nevidí)" in manager.get(
        "/calendar"
    ).get_data(as_text=True)


# ---------------------------------------------------------------------
# Práva
# ---------------------------------------------------------------------

def test_manager_cannot_publish_a_foreign_branch(pub_world):
    manager = login(pub_world, "mgr_a", "managerpass1")
    ids = pub_world.test_ids

    response = publish_action(manager, pub_world, "publish", branch="b")

    assert "Vyber pobočku" in response.get_data(as_text=True)
    assert published_flags(pub_world)[ids["s_boris"]] == 0


def test_employee_cannot_publish(pub_world):
    employee = login(pub_world, "emp_b", "emppass12")
    ids = pub_world.test_ids

    assert employee.get("/shifts/publish").status_code == 403

    token = _csrf(employee.get("/my-schedule").get_data(as_text=True))
    response = employee.post(
        "/shifts/publish",
        data={
            "csrf_token": token,
            "action": "publish",
            "branch_id": str(ids["b"]),
            "week_date": "2026-09-21",
        },
    )

    assert response.status_code == 403
    assert published_flags(pub_world)[ids["s_boris"]] == 0


def test_employee_with_manager_function_can_publish_and_sees_published_own(pub_world):
    from app.services.auth_service import set_manager_function

    ids = pub_world.test_ids

    with pub_world.app_context():
        set_manager_function(ids["anna"], True, [ids["a"]])

    worker_manager = login(pub_world, "emp_a", "emppass12")

    # Vlastný rozpis ("Môj rozpis") ukazuje len zverejnené smeny.
    html = worker_manager.get("/my-schedule").get_data(as_text=True)
    assert "2026-09-21" not in html

    publish_action(worker_manager, pub_world, "publish")

    assert "2026-09-21" in worker_manager.get("/my-schedule").get_data(as_text=True)


@pytest.mark.parametrize(
    "overrides, message",
    [
        ({"branch_id": ""}, "Vyber pobočku"),
        ({"week_date": "nie-datum"}, "platný dátum"),
    ],
)
def test_invalid_publish_input_changes_nothing(pub_world, overrides, message):
    admin = login(pub_world, "admin", "adminpass1")
    token = _csrf(admin.get("/shifts/publish").get_data(as_text=True))

    data = {
        "csrf_token": token,
        "action": "publish",
        "branch_id": str(pub_world.test_ids["a"]),
        "week_date": "2026-09-21",
    }
    data.update(overrides)

    response = admin.post("/shifts/publish", data=data)

    assert message in response.get_data(as_text=True)
    assert published_flags(pub_world)[pub_world.test_ids["s_mon"]] == 0


# ---------------------------------------------------------------------
# Kopírovanie týždňa a zverejnenie
# ---------------------------------------------------------------------

def _copy_week(client, app):
    token = _csrf(client.get("/shifts/copy-week").get_data(as_text=True))

    return client.post(
        "/shifts/copy-week",
        data={
            "csrf_token": token,
            "action": "copy",
            "branch_id": str(app.test_ids["a"]),
            "source_date": "2026-09-21",
            "target_date": "2026-10-05",
        },
    )


def test_copied_shifts_are_drafts_and_result_page_links_to_publishing(pub_world):
    manager = login(pub_world, "mgr_a", "managerpass1")

    html = _copy_week(manager, pub_world).get_data(as_text=True)

    assert "koncept" in html
    assert f"/shifts/publish?branch_id={pub_world.test_ids['a']}&amp;week_date=2026-10-05" in html

    from datetime import date

    from app.services.publish_service import week_publication_summary

    with pub_world.app_context():
        assert week_publication_summary(date(2026, 10, 5), pub_world.test_ids["a"]) == {
            "total": 2, "drafts": 2, "published": 0,
        }


def test_copied_shifts_are_published_with_auto_publish(pub_world, monkeypatch):
    monkeypatch.setenv("AUTO_PUBLISH_SHIFTS", "1")
    manager = login(pub_world, "mgr_a", "managerpass1")

    html = _copy_week(manager, pub_world).get_data(as_text=True)

    assert "/shifts/publish?" not in html

    from datetime import date

    from app.services.publish_service import week_publication_summary

    with pub_world.app_context():
        assert week_publication_summary(date(2026, 10, 5), pub_world.test_ids["a"]) == {
            "total": 2, "drafts": 0, "published": 2,
        }


# ---------------------------------------------------------------------
# Migrácia
# ---------------------------------------------------------------------

def test_migration_keeps_existing_shifts_published(tmp_path):
    """Doterajšie smeny zamestnanci už vidia - po migrácii ostávajú
    zverejnené. Smeny pridané potom sú koncept."""

    import sqlalchemy as sa
    from flask_migrate import downgrade, upgrade

    from app.extensions import db
    from app.web.app import create_app

    app = create_app(database_path=tmp_path / "pubmig.db")

    with app.app_context():
        upgrade(revision="c4d5e6f7a809")

        connection = db.session.connection()
        branch_id = connection.execute(sa.text("SELECT id FROM branches")).scalar()
        connection.execute(sa.text(
            "INSERT INTO employees (first_name, last_name, position, "
            "employment_type, weekly_hours, active, branch_id) "
            "VALUES ('Jan', 'Novak', 'Op', 'custom', 40, 1, :b)"
        ), {"b": branch_id})
        connection.execute(sa.text(
            "INSERT INTO shifts (employee_id, shift_date, start_time, "
            "end_time, shift_type) VALUES (1, '2026-09-21', '06:00', '14:00', 'Ranná')"
        ))
        db.session.commit()

        upgrade()

        connection = db.session.connection()
        assert connection.execute(sa.text("SELECT published FROM shifts")).scalars().all() == [1]

        # Nový riadok bez uvedenia príznaku je koncept (predvolená hodnota).
        connection.execute(sa.text(
            "INSERT INTO shifts (employee_id, shift_date, start_time, "
            "end_time, shift_type) VALUES (1, '2026-09-22', '06:00', '14:00', 'Ranná')"
        ))
        db.session.commit()

        connection = db.session.connection()
        assert connection.execute(
            sa.text("SELECT published FROM shifts ORDER BY id")
        ).scalars().all() == [1, 0]
        db.session.commit()

        downgrade(revision="c4d5e6f7a809")

        columns = {c["name"] for c in sa.inspect(db.engine).get_columns("shifts")}
        assert "published" not in columns
